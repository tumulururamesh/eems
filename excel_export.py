from openpyxl import Workbook

from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment
)

from openpyxl.utils import get_column_letter


# ===========================================================
# Common Excel Styles
# ===========================================================

HEADER_FILL = PatternFill(
    fill_type="solid",
    start_color="1F4E78",
    end_color="1F4E78"
)

COLUMN_FILL = PatternFill(
    fill_type="solid",
    start_color="4472C4",
    end_color="4472C4"
)

WHITE_FONT = Font(
    bold=True,
    color="FFFFFF"
)

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)

CENTER_ALIGN = Alignment(
    horizontal="center",
    vertical="center"
)


# ===========================================================
# Generic Excel Export Function
# ===========================================================

def export_report_to_excel(
    report_title,
    academic_year,
    headers,
    rows,
    totals,
    organisation_name="EKALAVYA FOUNDATION",
    system_name="Ekalavya Education Management System"
):

    wb = Workbook()

    ws = wb.active

    ws.title = "Report"

    # -------------------------------------------------------
    # Report Header
    # -------------------------------------------------------

    column_count = len(headers)

    last_column = get_column_letter(column_count)

    ws.merge_cells(f"A1:{last_column}1")
    ws.merge_cells(f"A2:{last_column}2")
    ws.merge_cells(f"A4:{last_column}4")
    ws.merge_cells(f"A5:{last_column}5")

    ws["A1"] = organisation_name
    ws["A2"] = system_name

    ws["A4"] = report_title
    ws["A5"] = f"Academic Year : {academic_year}"

    for cell in ["A1", "A2", "A4", "A5"]:

        ws[cell].fill = HEADER_FILL
        ws[cell].font = WHITE_FONT
        ws[cell].alignment = CENTER_ALIGN

    # -------------------------------------------------------
    # Column Headings
    # -------------------------------------------------------

    row_no = 7

    for col, header in enumerate(headers, start=1):

        cell = ws.cell(
            row=row_no,
            column=col,
            value=header
        )

        cell.font = WHITE_FONT
        cell.fill = COLUMN_FILL
        cell.border = THIN_BORDER
        cell.alignment = CENTER_ALIGN

    # -------------------------------------------------------
    # Data
    # -------------------------------------------------------

    row_no += 1

    for record in rows:

        for col, value in enumerate(record.values(), start=1):

            cell = ws.cell(
                row=row_no,
                column=col,
                value=value
            )

            cell.border = THIN_BORDER

        row_no += 1

    # -------------------------------------------------------
    # Report Summary
    # -------------------------------------------------------

    row_no += 2

    ws.cell(
        row=row_no,
        column=1,
        value="Report Summary"
    )

    ws.cell(
        row=row_no + 1,
        column=1,
        value=totals
    )

    # -------------------------------------------------------
    # Auto-fit Column Widths
    # -------------------------------------------------------

    for column_cells in ws.columns:

        length = 0

        column = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            if cell.value is not None:

                length = max(
                    length,
                    len(str(cell.value))
                )

        ws.column_dimensions[column].width = length + 3

    return wb